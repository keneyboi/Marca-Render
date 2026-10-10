import csv
import io
import openpyxl
from django.db import transaction
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.http import HttpResponse
from django.db.models import Q, Count
from django.core.paginator import Paginator
from django.utils.functional import cached_property
from django.contrib.auth.decorators import login_required
from openpyxl.utils import get_column_letter
from django.views.decorators.http import require_POST
from apps.core.models import Event, AttendanceRecord
from .forms import EventForm, EditEventForm
from django.core.exceptions import PermissionDenied
from apps.core.access import get_data_owner

def _clean_cell_value(val):
    if val is None:
        return ''
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    return str(val).strip()


def _normalize_attendee(row_dict, row_number):
    """
    Sanitizes row keys/values and handles both:
    1. Separate 'first_name' and 'last_name' columns
    2. A single 'full name' / 'name' column (e.g. 'LASTNAME, FIRSTNAME')
    Accepts aliases: 'id number' -> student_id, 'email address' -> email.
    """
    clean_row = {
        str(k).strip().lower(): _clean_cell_value(v)
        for k, v in row_dict.items() if k is not None
    }

    # Ignore trailing empty lines
    if not any(clean_row.values()):
        return None

    # Handle First and Last Name resolution
    first_name = clean_row.get('first_name', '')
    last_name = clean_row.get('last_name', '')

    if not first_name or not last_name:
        # Check for unified "full name" or "name" column
        full_name = clean_row.get('full name') or clean_row.get('name') or ''
        if full_name:
            if ',' in full_name:
                parts = full_name.split(',', 1)
                last_name = parts[0].strip().title()
                first_name = parts[1].strip().title()
            else:
                parts = full_name.split()
                if len(parts) >= 2:
                    first_name = ' '.join(parts[:-1]).title()
                    last_name = parts[-1].title()
                else:
                    first_name = full_name.strip().title()
                    last_name = ''

    # Handle Student ID aliases
    student_id = (
        clean_row.get('student_id') or 
        clean_row.get('id number') or 
        clean_row.get('id_number') or 
        clean_row.get('student id') or 
        None
    )

    # Handle Email aliases
    email = (
        clean_row.get('email') or 
        clean_row.get('email address') or 
        clean_row.get('email_address') or 
        None
    )

    # Validation guards
    if not first_name:
        raise ValueError(f"Row {row_number}: First name is missing.")
    if not last_name:
        raise ValueError(f"Row {row_number}: Last name is missing.")
    if not student_id and not email:
        raise ValueError(f"Row {row_number} ({first_name} {last_name}): Either Student ID or Email must be provided.")

    return {
        'first_name': first_name,
        'last_name': last_name,
        'student_id': student_id,
        'email': email,
        'course': clean_row.get('course') or None,
        'year_level': clean_row.get('year_level') or None,
    }


def parse_roster_file(uploaded_file):
    records_to_create = []
    seen_student_ids = set()
    seen_emails = set()
    filename = uploaded_file.name.lower()

    def is_header_row(cleaned_cells):
        has_separate_names = 'first_name' in cleaned_cells and 'last_name' in cleaned_cells
        has_full_name = 'full name' in cleaned_cells or 'name' in cleaned_cells
        has_id = any(h in cleaned_cells for h in ('student_id', 'id number', 'id_number', 'student id'))
        has_email = any(h in cleaned_cells for h in ('email', 'email address', 'email_address'))
        return (has_separate_names or has_full_name) and (has_id or has_email)

    if filename.endswith('.csv'):
        raw_content = uploaded_file.read()
        try:
            text = raw_content.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = raw_content.decode('latin-1')

        lines = text.splitlines()

        header_idx = -1
        for idx, line in enumerate(lines):
            row_items = [item.strip().lower() for item in line.split(',')]
            if is_header_row(row_items):
                header_idx = idx
                break

        if header_idx == -1:
            raise ValueError("Could not find valid column headers. Ensure the file has Name ('first_name'/'last_name' or 'full name') and an ID or Email.")

        reader = csv.DictReader(lines[header_idx:])
        for idx, row in enumerate(reader, start=header_idx + 2):
            attendee = _normalize_attendee(row, row_number=idx)
            if attendee is None:
                continue

            sid, em = attendee['student_id'], attendee['email']
            if sid:
                sid_lower = sid.lower()
                if sid_lower in seen_student_ids:
                    raise ValueError(f"Row {idx}: Duplicate Student ID '{sid}' found in file.")
                seen_student_ids.add(sid_lower)
            if em:
                em_lower = em.lower()
                if em_lower in seen_emails:
                    raise ValueError(f"Row {idx}: Duplicate Email '{em}' found in file.")
                seen_emails.add(em_lower)

            records_to_create.append(attendee)

    elif filename.endswith(('.xlsx', '.xlsm')):
        wb = openpyxl.load_workbook(uploaded_file, read_only=True, data_only=True)
        sheet = wb.active

        rows_iter = sheet.iter_rows(values_only=True)
        headers = None
        start_row_num = 0

        for idx, row in enumerate(rows_iter, start=1):
            cleaned = [str(cell).strip().lower() if cell is not None else '' for cell in row]
            if is_header_row(cleaned):
                headers = cleaned
                start_row_num = idx
                break

        if not headers:
            wb.close()
            raise ValueError("Could not find valid column headers. Ensure the sheet has Name ('first_name'/'last_name' or 'full name') and an ID or Email.")

        for idx, row in enumerate(rows_iter, start=start_row_num + 1):
            row_dict = {
                headers[c_idx]: val
                for c_idx, val in enumerate(row)
                if c_idx < len(headers)
            }
            attendee = _normalize_attendee(row_dict, row_number=idx)
            if attendee is None:
                continue

            sid, em = attendee['student_id'], attendee['email']
            if sid:
                sid_lower = sid.lower()
                if sid_lower in seen_student_ids:
                    wb.close()
                    raise ValueError(f"Row {idx}: Duplicate Student ID '{sid}' found in file.")
                seen_student_ids.add(sid_lower)
            if em:
                em_lower = em.lower()
                if em_lower in seen_emails:
                    wb.close()
                    raise ValueError(f"Row {idx}: Duplicate Email '{em}' found in file.")
                seen_emails.add(em_lower)

            records_to_create.append(attendee)

        wb.close()
    else:
        raise ValueError('Unsupported file format. Please upload a .csv or .xlsx file.')

    if not records_to_create:
        raise ValueError("Roster contains no valid attendee records.")

    return records_to_create


@login_required(login_url='landing')
def create_event(request):
    if request.method != 'POST':
        return redirect('home')

    form = EventForm(request.POST, request.FILES)

    if not form.is_valid():
        messages.error(request, 'Please check the event details for errors.')
        return redirect('home')

    uploaded_file = request.FILES.get('roster_file')
    if not uploaded_file:
        messages.error(request, 'Please upload a roster file.')
        return redirect('home')

    # Parse and validate the file BEFORE creating the event
    try:
        records_to_create = parse_roster_file(uploaded_file)
    except Exception as e:
        # Abort immediately if any error occurs
        messages.error(request, f'Roster error: {e}')
        return redirect('home')

    # Atomic creation: If any database error occurs, roll back the event completely
    try:
        with transaction.atomic():
            event = form.save(commit=False)
            event.user = request.user
            event.save()

            attendees = [
                AttendanceRecord(event=event, **data)
                for data in records_to_create
            ]
            AttendanceRecord.objects.bulk_create(attendees)

        messages.success(request, f'Event "{event.name}" and {len(records_to_create)} attendee records created successfully!')
        return redirect('home')

    except Exception as e:
        messages.error(request, f'Failed to create event: {e}')
        return redirect('home')


@login_required(login_url='landing')
def download_attendance_template_csv(request):
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="attendance_template.csv"'

    writer = csv.writer(response)
    writer.writerow(['student_id', 'email', 'first_name', 'last_name', 'course', 'year_level'])
    writer.writerow(['21-1234-567', '', 'John', 'Doe', 'BSCS', '3'])
    writer.writerow(['', 'jane.guest@company.com', 'Jane', 'Smith', '', ''])
    return response


@login_required(login_url='landing')
def download_attendance_template_xlsx(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance Template"

    headers = ['student_id', 'email', 'first_name', 'last_name', 'course', 'year_level']
    ws.append(headers)

    ws.append(['21-1234-567', '', 'John', 'Doe', 'BSCS', '3'])
    ws.append(['', 'jane.guest@company.com', 'Jane', 'Smith', '', ''])

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="attendance_template.xlsx"'
    wb.save(response)
    return response


class _KnownCountPaginator(Paginator):
    """Paginator that reuses a count we already computed (skips its COUNT query)."""

    def __init__(self, *args, known_count, **kwargs):
        super().__init__(*args, **kwargs)
        self._known_count = known_count

    @cached_property
    def count(self):
        return self._known_count


@login_required(login_url='landing')
def event_detail(request, event_id):
    owner_user = get_data_owner(request)
    if owner_user is None:
        raise PermissionDenied("You do not have permission to access these records.")

    event = get_object_or_404(Event, id=event_id, user=owner_user)
    records_qs = event.attendance_records
    all_records = records_qs.select_related('edited_by').order_by('last_name', 'first_name')

    # Distinct values only (a few rows) instead of loading the whole roster.
    courses = sorted({
        c.strip()
        for c in records_qs.order_by().values_list('course', flat=True).distinct()
        if c and c.strip() and c.strip().lower() != 'nan'
    })

    year_levels = sorted({
        str(y).strip()
        for y in records_qs.order_by().values_list('year_level', flat=True).distinct()
        if y is not None and str(y).strip() and str(y).strip().lower() != 'nan'
    }, key=lambda y: int(y) if y.isdigit() else y)

    # Totals, status counts and "has student IDs" in ONE query.
    stats = records_qs.aggregate(
        total=Count('id'),
        present=Count('id', filter=Q(status=AttendanceRecord.Status.PRESENT)),
        late=Count('id', filter=Q(status=AttendanceRecord.Status.LATE)),
        absent=Count('id', filter=Q(status=AttendanceRecord.Status.ABSENT)),
        with_student_id=Count(
            'id',
            filter=~Q(student_id__isnull=True) & ~Q(student_id='') & ~Q(student_id='nan'),
        ),
    )
    has_student_id = stats['with_student_id'] > 0
    has_academic_info = bool(courses or year_levels)

    paginator = _KnownCountPaginator(all_records, 50, known_count=stats['total'])
    page_number = request.GET.get('page')
    records = paginator.get_page(page_number)

    st = str(event.session_type).strip()
    show_session_1 = st == '1'
    show_session_2 = st == '2'
    show_session_3 = st == '3'

    context = {
        'event': event,
        'records': records,
        'status_choices': AttendanceRecord.Status.choices,
        'unique_courses': courses,
        'unique_year_levels': year_levels,
        'show_session_1': show_session_1,
        'show_session_2': show_session_2,
        'show_session_3': show_session_3,
        'has_student_id': has_student_id,
        'has_academic_info': has_academic_info,
        'total_attendance': stats['total'],
        'present_count': stats['present'],
        'late_count': stats['late'],
        'absent_count': stats['absent'],
    }
    return render(request, 'event/event_detail.html', context)


@login_required(login_url='landing')
@require_POST
def add_attendance_record(request, event_id):
    event = get_object_or_404(Event, id=event_id, user=request.user)

    student_id = request.POST.get('student_id', '').strip() or None
    email = request.POST.get('email', '').strip() or None
    first_name = request.POST.get('first_name', '').strip()
    last_name = request.POST.get('last_name', '').strip()
    course = request.POST.get('course', '').strip() or None
    year_level = request.POST.get('year_level', '').strip() or None
    status = request.POST.get('status', AttendanceRecord.Status.ABSENT)

    if not first_name or not last_name:
        messages.error(request, 'First name and Last name are required.')
        return redirect('event_detail', event_id=event.id)

    if not student_id and not email:
        messages.error(request, 'Either Student ID or Email must be provided.')
        return redirect('event_detail', event_id=event.id)

    if student_id and AttendanceRecord.objects.filter(event=event, student_id__iexact=student_id).exists():
        messages.error(request, f'An attendee with Student ID "{student_id}" already exists for this event.')
        return redirect('event_detail', event_id=event.id)

    if email and AttendanceRecord.objects.filter(event=event, email__iexact=email).exists():
        messages.error(request, f'An attendee with Email "{email}" already exists for this event.')
        return redirect('event_detail', event_id=event.id)

    AttendanceRecord.objects.create(
        event=event,
        email=email,
        first_name=first_name,
        last_name=last_name,
        student_id=student_id,
        course=course,
        year_level=year_level,
        status=status,
        edited_by=request.user,
    )

    messages.success(request, f'Attendee "{first_name} {last_name}" added successfully.')
    return redirect('event_detail', event_id=event.id)


@login_required(login_url='landing')
@require_POST
def delete_attendance_record(request, event_id, record_id):
    event = get_object_or_404(Event, id=event_id, user=request.user)
    record = get_object_or_404(AttendanceRecord, id=record_id, event=event)
    
    attendee_name = f"{record.first_name} {record.last_name}".strip()
    record.delete()
    
    messages.success(request, f'Attendee "{attendee_name}" was removed from attendance records.')
    return redirect('event_detail', event_id=event.id)


@login_required(login_url='landing')
@require_POST
def delete_event(request, event_id):
    event = get_object_or_404(Event, id=event_id, user=request.user)
    name = event.name
    event.delete()
    messages.success(request, f'Event "{name}" and its attendance was successfully deleted.')
    return redirect('home')


@login_required(login_url='landing')
@require_POST
def edit_event(request, event_id):
    event = get_object_or_404(Event, id=event_id, user=request.user)

    # Use EditEventForm instead of EventForm
    form = EditEventForm(request.POST, instance=event)

    if not form.is_valid():
        first_error = next(iter(form.errors.values()))[0] if form.errors else 'Invalid input.'
        messages.error(request, f'Failed to update event: {first_error}')
        return redirect('home')

    try:
        updated_event = form.save()
        messages.success(request, f'Event "{updated_event.name}" was successfully updated!')
    except Exception as e:
        messages.error(request, f'Failed to update event: {e}')

    return redirect('home')